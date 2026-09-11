import io
import json
from pathlib import Path
import re

from openpyxl import Workbook
import pytest

import gen


def test_funtides_uses_standard_software_sectioning():
    template = (gen.ROOT / "templates/desc-software.tex").read_text(encoding="utf-8")
    fiche = (gen.ROOT / "software/funtides/funtides.tex").read_text(encoding="utf-8")
    headings = r"\\subsection\{([^}]+)\}"
    expected = re.findall(headings, template)
    # This optional extension is also used by the FreeFEM++ and HPDDM fiches.
    expected.insert(expected.index("Mathematics"), "Application entry points")
    assert re.findall(headings, fiche) == expected
    # General information precedes the numbered subsections in the template.
    assert fiche.index(r"\input{software/funtides/metadata.tex}") < fiche.index(
        r"\subsection{Software summary}"
    )


@pytest.fixture
def snapshot():
    return json.loads((gen.ROOT / "software/funtides/catalogue.json").read_text())


@pytest.mark.parametrize("value", [False, "FALSE", "false", "No", 0, "0", 0.0])
def test_false_flags(value):
    assert gen.benchmark_flag(value) is False


@pytest.mark.parametrize("value", [True, "TRUE", "true", "Yes", 1, "1", 1.0])
def test_true_flags(value):
    assert gen.benchmark_flag(value) is True


@pytest.mark.parametrize("value", [None, "", "  "])
def test_missing_flags(value):
    assert gen.benchmark_flag(value) is None


def test_invalid_flag():
    with pytest.raises(ValueError, match="Invalid"):
        gen.benchmark_flag("pending")


def test_latex_escape():
    assert gen.latex(r"a_b & 50% #1 $x {y} \ ~ ^") == (
        r"a\_b \& 50\% \#1 \$x \{y\} \textbackslash{} "
        r"\textasciitilde{} \textasciicircum{}"
    )
    assert gen.latex("Aurélien\nC++17") == "Aurélien C++17"


def test_links():
    assert gen.link("https://example.org/a_b?q=1&x=2#part", "Docs") == (
        r"\href{https://example.org/a\_b?q=1\&x=2\#part}{Docs}"
    )
    with pytest.raises(ValueError, match="HTTP"):
        gen.link("file:///etc/passwd", "Bad")
    assert r"\%7B" in gen.link("https://example.org/{unsafe}", "Docs")


@pytest.mark.parametrize(
    "sheet,name", [("Frameworks", "Name"), ("Software", "Software")]
)
def test_xlsx_current_and_legacy(tmp_path, sheet, name):
    workbook = Workbook()
    workbook.active.title = sheet
    workbook.active.append([name, "WP1 Benchmarked", "Description"])
    workbook.active.append(["Other", True, "Do not select"])
    workbook.active.append(["FUnTiDES", False, "Éléments spectraux"])
    path = tmp_path / "catalogue.xlsx"
    workbook.save(path)
    result = gen.load_local(path, "funtides")
    assert result["framework"]["Name"] == "FUnTiDES"
    assert result["framework"]["WP1 Benchmarked"] is False
    assert result["source"]["row"] == 3
    assert result["source"]["sheet"] == sheet


def test_sheet_precedence_and_override(tmp_path):
    workbook = Workbook()
    workbook.active.title = "Software"
    workbook.active.append(["Software"])
    workbook.active.append(["Old"])
    sheet = workbook.create_sheet("Frameworks")
    sheet.append(["Name"])
    sheet.append(["FUnTiDES"])
    path = tmp_path / "catalogue.xlsx"
    workbook.save(path)
    assert gen.load_local(path, "FUnTiDES")["source"]["sheet"] == "Frameworks"
    assert gen.load_local(path, "Old", "Software")["framework"]["Name"] == "Old"
    with pytest.raises(ValueError, match="worksheet"):
        gen.load_local(path, "FUnTiDES", "Missing")


def test_csv(tmp_path):
    path = tmp_path / "catalogue.csv"
    path.write_text(
        '\ufeffName,WP1 Benchmarked,Description\nFUnTiDES,FALSE,"CPU, GPU"\n'
    )
    result = gen.load_local(path, "FUnTiDES")
    assert result["framework"]["Description"] == "CPU, GPU"
    assert gen.benchmark_flag(result["framework"]["WP1 Benchmarked"]) is False


@pytest.mark.parametrize(
    "rows,error",
    [
        ([["Name"], ["Other"]], "found 0"),
        ([["Name"], ["FUnTiDES"], ["funtides"]], "found 2"),
        ([["Name", "Name"], ["FUnTiDES", "Other"]], "Duplicate"),
        ([["NotName"], ["FUnTiDES"]], "Expected a Name"),
    ],
)
def test_invalid_catalogues(rows, error):
    with pytest.raises(ValueError, match=error):
        gen.select_row(rows, "FUnTiDES", {})


def test_read_only_google_export(monkeypatch):
    calls = []

    class Response(io.BytesIO):
        headers = {"Content-Type": "text/csv"}

    def fake_open(url, timeout):
        calls.append((url, timeout))
        return Response(b"Name,WP1 Benchmarked\nFUnTiDES,FALSE\n")

    monkeypatch.setattr(gen, "urlopen", fake_open)
    result = gen.load_google(
        "https://docs.google.com/spreadsheets/d/test_id/edit?gid=1#gid=42",
        "FUnTiDES",
        "Frameworks",
    )
    assert calls == [
        ("https://docs.google.com/spreadsheets/d/test_id/export?format=csv&gid=42", 30)
    ]
    assert result["source"]["row"] == 2
    assert result["source"]["gid"] == "42"


def test_google_login_rejected(monkeypatch):
    class Response(io.BytesIO):
        headers = {"Content-Type": "text/html"}

    monkeypatch.setattr(
        gen, "urlopen", lambda *a, **k: Response(b"<html>Sign in</html>")
    )
    with pytest.raises(ValueError, match="authentication"):
        gen.load_google("https://docs.google.com/spreadsheets/d/test/edit", "FUnTiDES")
    with pytest.raises(ValueError, match="docs.google.com"):
        gen.load_google("https://example.org/test", "FUnTiDES")


def test_generated_snapshot_is_current(snapshot):
    expected = gen.ROOT / "software/funtides/metadata.tex"
    assert expected.read_text() == gen.render(snapshot)
    text = gen.render(snapshot)
    assert "No WP benchmark flag is checked" in text
    assert "Flags not supplied" not in text
    assert "WP7 topics & Geophysics" in text
    assert "None" not in text
    assert "provide short description" not in text


def test_missing_is_not_false(snapshot):
    snapshot["framework"]["WP1 Benchmarked"] = ""
    assert "Flags not supplied: WP1." in gen.render(snapshot)
    snapshot["framework"]["WP1 Benchmarked"] = True
    assert "Benchmark flags checked in the survey: WP1." in gen.render(snapshot)


def test_targeted_generation_and_check(tmp_path, snapshot, monkeypatch):
    monkeypatch.chdir(tmp_path)
    source = tmp_path / "catalogue.json"
    source.write_text(json.dumps(snapshot))
    editorial = tmp_path / "funtides.tex"
    editorial.write_text("Handwritten scientific text\n")
    output = tmp_path / "metadata.tex"
    args = ["--input", str(source), "--framework", "FUnTiDES", "--output", str(output)]
    assert gen.main(args) == 0
    assert gen.main(args + ["--check"]) == 0
    assert gen.main(args) == 0
    assert editorial.read_text() == "Handwritten scientific text\n"
    assert set(p.name for p in tmp_path.iterdir()) == {
        "catalogue.json",
        "metadata.tex",
        "funtides.tex",
    }
    output.write_text(gen.MARKER + "Stale\n")
    with pytest.raises(SystemExit) as error:
        gen.main(args + ["--check"])
    assert error.value.code == 2
    assert output.read_text() == gen.MARKER + "Stale\n"


def test_no_overwrite_editorial_or_snapshot(tmp_path, snapshot):
    source = tmp_path / "catalogue.json"
    source.write_text(json.dumps(snapshot))
    output = tmp_path / "funtides.tex"
    output.write_text("Handwritten scientific text\n")
    args = ["--input", str(source), "--framework", "FUnTiDES", "--output", str(output)]
    with pytest.raises(SystemExit):
        gen.main(args)
    assert output.read_text() == "Handwritten scientific text\n"
    with pytest.raises(SystemExit):
        gen.main(args[:-1] + [str(source)])
    with pytest.raises(SystemExit):
        gen.main(args + ["--snapshot", str(source)])
    with pytest.raises(SystemExit):
        gen.main(["--input", str(source), "--framework", "Other"])


def test_save_selected_snapshot(tmp_path):
    source = tmp_path / "catalogue.csv"
    source.write_text(
        "Name,Emails\nOther,other@example.org\nFUnTiDES,contact@example.org\n"
    )
    output = tmp_path / "metadata.tex"
    snapshot = tmp_path / "selected.json"
    args = [
        "--input",
        str(source),
        "--framework",
        "FUnTiDES",
        "--output",
        str(output),
        "--snapshot",
        str(snapshot),
    ]
    assert gen.main(args) == 0
    assert "other@example.org" not in snapshot.read_text()
    assert json.loads(snapshot.read_text())["source"]["row"] == 3
    with pytest.raises(SystemExit):
        gen.main(args)
